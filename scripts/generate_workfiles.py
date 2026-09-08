"""Generate AVD specialization workfiles (docx/pptx/xlsx) into static/templates/."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pptx import Presentation
from pptx.util import Inches as PInches
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
PUB = ROOT / "static" / "templates"
ENG = PUB / "engagement"
DEL = PUB / "deliverables"
AUD = PUB / "audit"
for p in (ENG, DEL, AUD):
    p.mkdir(parents=True, exist_ok=True)

BLUE = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x59, 0x59, 0x59)


def docx_new(title: str, subtitle: str) -> Document:
    d = Document()
    style = d.styles["Normal"]; style.font.name = "Calibri"; style.font.size = Pt(11)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("\n\n\nAzure Virtual Desktop\nMicrosoft Advanced Specialization")
    r.font.size = Pt(14); r.font.color.rgb = GREY
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("\n" + title); r.font.size = Pt(28); r.font.bold = True; r.font.color.rgb = BLUE
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle); r.font.size = Pt(13); r.font.color.rgb = GREY
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("\n\nPartner template — customise per customer engagement\n\nVersion 1.0")
    r.font.size = Pt(10); r.font.color.rgb = GREY
    d.add_page_break()
    d.add_heading("Table of Contents", level=1)
    note = d.add_paragraph("Right-click and Update Field in Word to populate.")
    note.runs[0].italic = True
    paragraph = d.add_paragraph(); run = paragraph.add_run()
    f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'), 'separate')
    txt = OxmlElement('w:t'); txt.text = "Right-click to update."
    f3 = OxmlElement('w:fldChar'); f3.set(qn('w:fldCharType'), 'end')
    rel = run._r
    for el in (f1, instr, f2, txt, f3):
        rel.append(el)
    d.add_page_break()
    return d


def add_section(doc, heading, prompts):
    doc.add_heading(heading, level=1)
    for prompt in prompts:
        p = doc.add_paragraph(); r = p.add_run(prompt); r.italic = True; r.font.color.rgb = GREY
        doc.add_paragraph("")


def pptx_new(title, subtitle):
    prs = Presentation()
    prs.slide_width = PInches(13.333); prs.slide_height = PInches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title; slide.placeholders[1].text = subtitle
    return prs


def add_slide(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    tf = slide.placeholders[1].text_frame
    tf.text = bullets[0] if bullets else ""
    for b in bullets[1:]:
        p = tf.add_paragraph(); p.text = b


def add_section_slide(prs, title):
    layout = prs.slide_layouts[2] if len(prs.slide_layouts) > 2 else prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title


HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(bold=True, color="FFFFFF", size=11)
STATUS_LIST = '"Not started,In progress,Blocked,Complete,N/A"'


def style_header(ws, headers, widths=None):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = HDR_FILL; c.font = HDR_FONT
        c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"; ws.row_dimensions[1].height = 28
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def add_status_dv(ws, col_letter, last_row=200):
    dv = DataValidation(type="list", formula1=STATUS_LIST, allow_blank=True)
    dv.add(f"{col_letter}2:{col_letter}{last_row}")
    ws.add_data_validation(dv)


def gen_offering_pptx():
    prs = pptx_new("Azure Virtual Desktop", "Partner Offering — One-pager")
    add_slide(prs, "Why AVD now",
        ["Hybrid work is permanent; secure remote desktop is a strategic capability",
         "Consolidate RDS / Citrix / Horizon on Azure-native PaaS",
         "Tight integration with Entra ID, Intune, Microsoft 365, Defender",
         "Predictable consumption with autoscale and Reserved Instances"])
    add_slide(prs, "What we deliver",
        ["Discovery workshop and persona analysis",
         "Reference architecture selection (pooled multi-session / personal / single-session)",
         "Image management pipeline with Azure Image Builder",
         "FSLogix profile design on Azure Files / ANF",
         "Identity, network and security hardening",
         "Production deployment via IaC, KT and hypercare"])
    add_slide(prs, "Engagement shape",
        ["Discovery: 1–2 weeks", "Design: 2–3 weeks", "Pilot: 2–4 weeks",
         "Production rollout: wave-based, 4–12 weeks", "Hypercare: 2–4 weeks with explicit exit criteria"])
    add_slide(prs, "Outcomes",
        ["Production AVD environment aligned to WAF", "Customer ops team confident to operate (KT signed off)",
         "Documented run state (HLD, LLD, runbook, BCDR)", "Defined cost baseline and optimisation levers"])
    add_slide(prs, "Why us",
        ["Microsoft Advanced Specialization for AVD", "Repeatable IaC and image pipelines",
         "Innersourced playbook updated every engagement", "Migration expertise from RDS, Citrix-on-Azure, Horizon-on-Azure"])
    add_slide(prs, "Next steps",
        ["Schedule a 60-minute discovery call", "Share access to the qualification questionnaire",
         "Agree scope and propose a fixed-price Discovery + Design"])
    prs.save(ENG / "offering-one-pager.pptx")


def gen_qual_docx():
    d = docx_new("Qualification Questionnaire", "Pre-engagement input for AVD opportunities")
    sections = [
        ("1. Business context", ["Driver: RDS EoL, Citrix renewal, M&A, security, cost, hybrid work, geographic expansion, regulated workload.",
            "Business outcome with measurable target.", "Executive sponsor and their success metric."]),
        ("2. User personas and app inventory", ["Per persona: headcount per region, workload profile, apps, peripherals, data sensitivity.",
            "Current end-user device estate.", "Current users on RDS / Citrix / Horizon / W365 with counts and pain points."]),
        ("3. Image strategy", ["Marketplace + customisations or full custom gold image?",
            "Apps in base image vs Intune / AppAttach / MSIX.", "Steady-state image owner.", "Patching cadence."]),
        ("4. Identity", ["Cloud-only Entra, hybrid, AD DS only.", "MFA coverage today.", "Conditional Access and device compliance signals.", "RBAC / PIM standards."]),
        ("5. Networking", ["Hub-spoke topology; ExpressRoute / VPN.", "DNS resolution requirements.", "Egress filtering pattern.",
            "RDP Shortpath for managed networks feasible?", "Data residency constraints."]),
        ("6. Storage and profiles", ["Azure Files Premium or ANF.", "Per-user profile size budget.", "Office container split required?", "RPO / RTO target for profiles."]),
        ("7. Security, compliance, governance", ["Regulatory regime.", "Data classification scheme.", "Endpoint isolation requirements.",
            "Defender for Cloud baseline.", "Azure Policy initiatives."]),
        ("8. Operations and support", ["Ops team size and AZ-140 / AZ-104 coverage.", "ITSM tool.", "Monitoring tooling.", "On-call coverage."]),
        ("9. Migration in scope", ["Source platform.", "Profile migration approach.", "Cutover model.", "Data migration volume and timeline."]),
        ("10. Commercial", ["Funding source.", "Available Microsoft 365 licensing.", "Procurement constraints.", "Decision timeline."]),
        ("11. Decision summary", ["Recommended engagement shape.", "Audit-fit assessment.", "Top three risks.", "Go / no-go and rationale."]),
    ]
    for h, prompts in sections:
        add_section(d, h, prompts)
    d.save(ENG / "qualification-questionnaire.docx")


def gen_discovery_pptx():
    prs = pptx_new("AVD Discovery Workshop", "Joint workshop — partner + customer EUC, identity, network, security")
    add_slide(prs, "Agenda — Day 1",
        ["Business drivers and target outcomes", "Persona walk-through", "Application inventory walk-through", "Image strategy options"])
    add_slide(prs, "Agenda — Day 2",
        ["Identity, network and security architecture review", "Profile storage and BCDR",
         "Reference architecture selection", "Risks, assumptions, dependencies and next steps"])
    for sect, bullets in [
        ("Business drivers", ["Why AVD, why now", "Measurable target outcomes", "Sponsor and decision structure", "Constraints"]),
        ("Personas", ["Per region: headcount, profile, peripherals, data sensitivity", "Current device estate", "Persona-to-architecture mapping draft"]),
        ("Application strategy", ["Apps in base image vs Intune vs AppAttach vs streamed", "M365 Apps shared-computer activation", "Teams optimised media path", "LOB and legacy thick-client handling"]),
        ("Image strategy", ["Marketplace + customisation vs custom gold image", "Azure Image Builder pipeline outline", "Validation ring / production ring", "Refresh cadence and owner"]),
        ("Identity, network, security", ["Entra-join vs hybrid", "Conditional Access scope and signals", "Network topology and DNS", "RDP Shortpath, egress filtering, private endpoints", "Defender for Cloud and Policy"]),
        ("Profile storage & BCDR", ["Azure Files Premium vs ANF", "FSLogix container topology", "Per-user budget and Office container split", "BCDR pattern and RPO / RTO"]),
        ("Next steps", ["Confirm persona-to-architecture mapping", "Sign-off discovery workbook", "Schedule HLD review", "Open WAF assessment"]),
    ]:
        add_section_slide(prs, sect); add_slide(prs, sect, bullets)
    prs.save(ENG / "discovery-workshop-deck.pptx")


def gen_discovery_xlsx():
    wb = Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("Personas")
    style_header(ws, ["Persona","Region","Headcount","Profile type","Apps","Peripherals","Data sensitivity","Current platform","Target architecture","Notes","Status"], [22,12,10,18,30,22,18,18,24,32,14])
    add_status_dv(ws, "K")
    ws = wb.create_sheet("Applications")
    style_header(ws, ["App name","Vendor","Delivery (Image / Intune / AppAttach / MSIX / Streamed)","Per-user vs shared","License model","Compatibility tested?","Owner","Notes","Status"], [24,16,38,18,16,18,18,32,14])
    add_status_dv(ws, "I")
    for sheet, topics in [
        ("Identity", ["Entra join model","MFA coverage","Conditional Access for AVD","Device compliance signal","RBAC standard","PIM eligible assignments","Privileged access workstation"]),
        ("Network", ["Hub-spoke topology","ExpressRoute / VPN","DNS resolution","Egress filtering","RDP Shortpath","Private endpoints","Data residency"]),
        ("Profiles & storage", ["Profile platform (Azure Files / ANF)","Authentication mode","Per-user GB budget","Office container split","Backup / restore","BCDR replication","RPO / RTO"]),
    ]:
        ws = wb.create_sheet(sheet)
        style_header(ws, ["Topic","Current state","Target state","Gap / action","Owner","Due","Status"], [28,28,28,32,16,12,14])
        for t in topics:
            ws.append([t])
        add_status_dv(ws, "G")
    ws = wb.create_sheet("Risks")
    style_header(ws, ["ID","Risk","Likelihood (L/M/H)","Impact (L/M/H)","Mitigation","Owner","Status"], [8,40,18,16,40,16,14])
    add_status_dv(ws, "G")
    ws = wb.create_sheet("Decisions")
    style_header(ws, ["ID","Decision","Options considered","Decision rationale","Owner","Date","Status"], [8,32,32,32,16,12,14])
    add_status_dv(ws, "G")
    wb.save(ENG / "discovery-workshop-workbook.xlsx")


def gen_waf_xlsx():
    wb = Workbook(); wb.remove(wb.active)
    pillars = {
        "Reliability": ["Paired-region host pool documented","FSLogix replication mechanism in place and tested","RTO / RPO targets agreed and tested","Autoscale plan covers peak and surge","Session host health alerts wired to action group"],
        "Security": ["Conditional Access policies scoped to AVD","MFA enforced for AVD users and admins","RBAC / PIM model documented and enforced","Defender for Cloud Servers Plan on session hosts","Private endpoints for storage / Key Vault / Log Analytics","Endpoint isolation policy signed off"],
        "Cost Optimization": ["Autoscale schedule reflects working hours","Reserved Instances / Savings Plan committed for steady-state","Storage tier matches profile IOPS need","Cost anomaly alerts active","Per-persona unit cost tracked"],
        "Operational Excellence": ["IaC pipeline owns production deployment","AVD Insights workbook in use","Image refresh runbook in place","FSLogix repair runbook in place","Quarterly operations review scheduled"],
        "Performance Efficiency": ["Session-density target measured per host SKU","Logon time P95 measured against target","RDP Shortpath used where feasible","Teams media optimised path verified","Profile container size monitored"],
    }
    ws = wb.create_sheet("Summary")
    style_header(ws, ["Pillar","Items","Complete","In progress","Blocked","Not started","Score"], [22,8,10,12,10,12,10])
    for i, name in enumerate(pillars.keys(), 2):
        ws.cell(row=i, column=1, value=name); ws.cell(row=i, column=2, value=len(pillars[name]))
    for name, rows in pillars.items():
        ws = wb.create_sheet(name)
        style_header(ws, ["#","Recommendation","Current state","Target state","Owner","Due","Status","Evidence link"], [6,42,32,32,16,12,14,28])
        for i, r in enumerate(rows, 1):
            ws.cell(row=i+1, column=1, value=i); ws.cell(row=i+1, column=2, value=r)
        add_status_dv(ws, "G")
    wb.save(ENG / "waf-assessment.xlsx")


def gen_assess_xlsx():
    wb = Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("MAP toolkit")
    style_header(ws, ["Input / report","Source","Purpose","Owner","Due","Status","Notes"], [32,18,32,16,12,14,32])
    for row in ["Hardware inventory CSV","Application inventory CSV","Session usage stats","Profile size distribution","Peak concurrent users","Migration wave proposal"]:
        ws.append([row])
    add_status_dv(ws, "F")
    for sheet, pairs in [
        ("RDS to AVD", [("RD Connection Broker","AVD service"),("RD Gateway","AVD service / RDP Shortpath"),("RD Web Access","AVD client / web client"),("RD Session Host","Session host VMs"),("RD Licensing","Windows 11/10 Enterprise multi-session + M365"),("User Profile Disks","FSLogix on Azure Files / ANF"),("Published RemoteApp","RemoteApp on AVD"),("GPOs","GPO + Intune + AVD per-host RDP properties")]),
        ("Citrix on Azure to AVD", [("Citrix Cloud Connector","AVD service"),("StoreFront","AVD client / web client"),("Delivery Controller","AVD service"),("VDAs (Server OS)","Pooled multi-session host pool"),("VDAs (Desktop OS)","Personal or single-session pooled"),("Citrix Profile Management","FSLogix"),("HDX policies","RDP properties + Intune CSP"),("App Layering","AppAttach / Intune"),("NetScaler Gateway","AVD service / RDP Shortpath")]),
        ("Horizon to AVD", [("Connection Server","AVD service"),("Horizon Client","AVD client"),("Instant Clones","Pooled multi-session + Image Builder"),("App Volumes","AppAttach / MSIX"),("DEM","FSLogix + GPO + Intune"),("Unified Access Gateway","AVD service / RDP Shortpath")]),
    ]:
        ws = wb.create_sheet(sheet)
        style_header(ws, ["Source component","Equivalent in AVD","Migration approach","Effort (S/M/L)","Owner","Status","Notes"], [28,28,32,14,16,14,32])
        for s, t in pairs:
            ws.append([s, t])
        add_status_dv(ws, "F")
    ws = wb.create_sheet("W365 interop")
    style_header(ws, ["Scenario","Recommendation","Notes","Status"], [32,32,40,14])
    for s in ["Knowledge worker, 1:1, predictable usage","Knowledge worker, 1:1, variable usage","Power user, GPU","Shared / pooled multi-session","Frontline shift worker","Contractor / temp"]:
        ws.append([s])
    add_status_dv(ws, "D")
    wb.save(ENG / "assessment-platform-inputs.xlsx")


def gen_dod_docx():
    d = docx_new("Definition of Done", "Engagement-level acceptance for an AVD production delivery")
    groups = [
        ("Pre Go-Live", ["[A.1.1] Business strategy signed by customer sponsor","[A.1.2] Adoption plan / wave plan agreed and tracked","[A.2.1] Security and governance tooling active","[A.2.2 + B.2.2] WAF review complete; register triaged","[A.3.1] IaC pipeline used for production","[B.1.1] Workload assessment signed off","[B.2.1] HLD signed by customer","[B.2.3] PoC / Pilot signed off with decision-to-proceed","[B.3.1] Cutover plan agreed (migrations)","[B.4.1] Test plan executed; results within targets"]),
        ("Go-Live", ["[B.3.1] Production pipeline run complete","[B.3.1] Configuration baseline captured","[B.3.1] Go-Live record signed","Communication sent to users","Hypercare started"]),
        ("Hypercare exit", ["All P1 / P2 incidents resolved","Incident rate below 1 / user / week (rolling 7 days)","Logon P95 within HLD target","Ops team confidence ≥ 4/5","Runbook validated"]),
        ("Handover", ["[B.4.2] LLD delivered","[B.4.2] Runbook delivered and walked through","[A.3.2 + B.4.2] KT plan complete; form signed","[B.4.2] Hypercare exit signed","[A.3.3] Ops tooling demonstrated","All evidence filed in tracker","Lessons learned filed"]),
        ("Audit readiness", ["All Module B evidence collected","Customer added to Evidence Tracker","Customer Case Study + Sign-Off captured","Ready as qualifying customer (if applicable)"]),
    ]
    for h, items in groups:
        d.add_heading(h, level=1)
        for it in items:
            d.add_paragraph(it, style="List Bullet")
    d.save(ENG / "definition-of-done.docx")


def gen_doc_template(filename, title, sections):
    d = docx_new(title, "Azure Virtual Desktop — template")
    for h, prompts in sections:
        add_section(d, h, prompts)
    d.save(DEL / filename)


def gen_hld_docx():
    gen_doc_template("hld-template.docx", "High-Level Design", [
        ("1. Executive summary", ["State purpose, scope and target outcome in three paragraphs."]),
        ("2. Business context", ["Link drivers and outcomes to A.1.1 strategy."]),
        ("3. Scope", ["In-scope personas, locations, regulatory regime. Out-of-scope items explicitly listed."]),
        ("4. Personas", ["Per persona: headcount, profile, peripherals, data sensitivity, target architecture."]),
        ("5. Reference architecture selection", ["For each persona: pattern, why, trade-offs considered."]),
        ("6. Host pool topology", ["For each pool: pattern, OS image, SKU, host count (baseline + max), max sessions per host, load balancing."]),
        ("7. Identity design", ["Entra-join vs hybrid; MFA; Conditional Access; RBAC / PIM model; PAW pattern."]),
        ("8. Profile storage design", ["Azure Files Premium vs ANF; auth; per-user GB; Office container split."]),
        ("9. Networking design", ["VNet, hub-spoke, ExpressRoute / VPN, RDP Shortpath, private endpoints, DNS, egress firewall."]),
        ("10. Image management strategy", ["Source, build pipeline, customisation steps, validation, promotion rings, refresh cadence."]),
        ("11. Application delivery strategy", ["Decision matrix per app; AppAttach storage and assignment; per-user vs shared activation."]),
        ("12. Scaling plan", ["Schedule, ramp-up safety, min / max, force-logoff behaviour, cost expectation."]),
        ("13. Monitoring & operations", ["Log Analytics, diagnostic settings, AVD Insights workbook, alert rules."]),
        ("14. Security & governance", ["Defender for Cloud, Policy, CA, endpoint isolation, data residency, log retention."]),
        ("15. BCDR pattern", ["Paired-region pattern, FSLogix replication, recovery, RPO / RTO, test cadence."]),
        ("16. Sizing & cost model", ["Sessions per host per persona, host SKU, totals, RI / SP plan, monthly forecast."]),
        ("17. Risks, assumptions, dependencies", ["Numbered RAID table; owner per item."]),
        ("18. Sign-off", ["Sponsor, customer EUC lead, security reviewer, partner engagement lead, partner technical lead."]),
    ])


def gen_lld_docx():
    gen_doc_template("lld-template.docx", "Low-Level Design (As-Built)", [
        ("1. Scope & version", ["Go-Live date; image version; IaC commit SHA."]),
        ("2. Subscription & MG structure", ["Tenant, MG, subscription, RG layout with names."]),
        ("3. Networking", ["VNets, subnets, NSGs, firewall rules, DNS zones, private endpoints."]),
        ("4. Host pools", ["Per pool: name, type, max session limit, load balancing, SKU, count, image reference."]),
        ("5. Session hosts", ["Current host names, deployment date, image version, tags."]),
        ("6. FSLogix profile storage", ["Storage account, share, permissions, container configuration."]),
        ("7. Image gallery", ["Gallery name, image definitions, current version, version retention."]),
        ("8. Identity", ["Entra groups, RBAC assignments, PIM eligible, CA policy IDs and conditions."]),
        ("9. Monitoring", ["Log Analytics workspace, diagnostic settings, AVD Insights workbook, alert rules, action groups."]),
        ("10. Scaling plans", ["Name, schedule, ramp-up / down configuration."]),
        ("11. App delivery", ["AppAttach packages, Intune app assignments, M365 Apps config."]),
        ("12. BCDR resources", ["Paired-region resources, replication mechanism, runbook references."]),
        ("13. IaC repository reference", ["Repo URL, commit SHA, parameter files."]),
        ("14. Change history since Go-Live", ["Date, change, ticket reference, owner."]),
    ])


def gen_runbook_docx():
    gen_doc_template("runbook-template.docx", "Operational Runbook", [
        ("1. Image refresh", ["Cadence; build via Azure Image Builder; validation pool; sign-off; rollout; rollback."]),
        ("2. Scaling", ["View current behaviour; adjust min / max; manage schedules; troubleshoot; cost impact."]),
        ("3. FSLogix", ["Container reset; container repair; share permissions audit; container growth review; Office container split."]),
        ("4. User onboarding / offboarding", ["Add to AVD group; pre-stage container; leaver archive / removal."]),
        ("5. Common incidents", ["Connection failure; slow logon; app launch failure; print failure; peripheral redirection."]),
        ("6. Patching", ["Ring-based update flow; WUfB / WSUS / Intune; image vs live patching; emergency patch."]),
        ("7. BCDR", ["Fail over to paired-region; fail back; FSLogix replication status; user communication."]),
        ("8. Backup / restore", ["What is backed up; RTO / RPO; restore procedures."]),
        ("9. Cost monitoring", ["Cost Management views; anomaly alerts; common levers."]),
        ("10. Escalation", ["L1 / L2 / L3 ownership; Microsoft support engagement; vendor escalation."]),
    ])


def gen_kt_docx():
    gen_doc_template("kt-plan-template.docx", "Knowledge Transfer Plan", [
        ("1. Audience and roles", ["EUC engineer, service desk lead, identity admin, operations manager, security reviewer."]),
        ("2. Session plan", ["Architecture, identity, network, FSLogix, image, scaling, incidents, runbook, shadowing, handover."]),
        ("3. Hands-on labs", ["Build a new gold image; reset a test user's FSLogix container; trigger and review an alert; add a new app via AppAttach."]),
        ("4. Evidence captured", ["Meeting invites; Teams attendance; recordings; lab completion; runbook walkthrough sign-off."]),
        ("5. KT exit criteria", ["All sessions delivered with named-role attendance; labs completed; runbook acknowledged; KT completion form signed."]),
        ("6. KT completion form", ["Sessions delivered; recordings handed over; labs completed; runbook acknowledged; readiness signed."]),
    ])


def gen_hyper_docx():
    gen_doc_template("hypercare-plan-template.docx", "Hypercare Plan", [
        ("1. Scope and duration", ["Go-Live to exit criteria; target 2–4 weeks; coverage hours; on-call P1 / P2 outside hours."]),
        ("2. Communication channels", ["Teams hypercare channel; email DL; on-call rota; daily stand-up; weekly review."]),
        ("3. Escalation matrix", ["P1 / P2 / P3 / P4 — definition, response time, owner."]),
        ("4. Daily cadence", ["09:00 stand-up; ticket review; on-call handover."]),
        ("5. Weekly cadence", ["Incident trends; AVD Insights; cost actual vs forecast; exit criteria progress."]),
        ("6. Exit criteria", ["All P1 / P2 resolved; incident rate target met; logon P95 within target; ops confidence ≥ 4/5; runbook validated."]),
        ("7. Exit sign-off", ["Customer ops manager; customer EUC lead; partner engagement lead."]),
        ("8. Post-hypercare arrangements", ["Steady-state support contract; lessons learned filed; next WAF re-assessment scheduled."]),
    ])


def gen_evidence_xlsx():
    wb = Workbook(); wb.remove(wb.active)
    a = [("A.1.1","Cloud & AI Adoption Business Strategy"),("A.1.2","Cloud & AI Adoption Plan"),("A.2.1","Security & Governance Tooling"),("A.2.2","Well-Architected Workloads"),("A.3.1","Repeatable Deployment"),("A.3.2","Plan for Skilling"),("A.3.3","Operations Management Tooling")]
    b = [("B.1.1","Workload Assessment"),("B.2.1","Solution Design"),("B.2.2","Azure Well-Architected Review"),("B.2.3","PoC or Pilot"),("B.3.1","Deployment to Production"),("B.4.1","Service Validation & Testing"),("B.4.2","Post-deployment Documentation")]
    for sheet, rows in [("Module A", a), ("Module B", b)]:
        ws = wb.create_sheet(sheet)
        style_header(ws, ["Control","Title","Customer / Tenant","Evidence item","Location / link","Owner","Due","Status","Auditor notes"], [10,36,22,40,32,16,12,14,32])
        for c, t in rows:
            ws.append([c, t])
        add_status_dv(ws, "H")
    ws = wb.create_sheet("Customers")
    style_header(ws, ["Customer","Engagement window","Module B controls covered","Sign-off received","Case study filed","Audit candidate (Y/N)","Status"], [22,22,28,18,16,18,14])
    add_status_dv(ws, "G")
    ws = wb.create_sheet("Audit cycle")
    style_header(ws, ["Item","Owner","Due","Status","Notes"], [40,16,14,14,32])
    for it in ["Confirm 2 qualifying customer engagements","Schedule audit with ISSI","Pay audit fee","Compile evidence pack","Run internal dry-run","Audit day","Receive certification","Renewal planning (year +1, +2)"]:
        ws.append([it])
    add_status_dv(ws, "D")
    wb.save(AUD / "evidence-tracker.xlsx")


def gen_prequal_xlsx():
    wb = Workbook(); wb.remove(wb.active)
    ws = wb.create_sheet("Pre-qualification")
    style_header(ws, ["#","Requirement","Source","Current state","Gap / action","Owner","Due","Status"], [6,42,18,28,32,16,12,14])
    items = [
        "Gold MS Partner or Solutions Partner — Modern Work / Infrastructure designation",
        "MPN / Partner ID active and in good standing",
        "Minimum AZ-140 certified individuals on staff (confirm current threshold from official guide)",
        "Minimum AZ-104 certified individuals on staff (confirm current threshold from official guide)",
        "Two qualifying customer engagements within the last 12 months",
        "Each qualifying engagement: monthly active user threshold met (confirm threshold from official guide)",
        "Customer sign-off letters for each qualifying engagement",
        "Customer case studies anonymisable for audit reference",
        "Audit fee budget approved",
        "Audit window booked with ISSI",
    ]
    for i, it in enumerate(items, 1):
        ws.cell(row=i+1, column=1, value=i)
        ws.cell(row=i+1, column=2, value=it)
    add_status_dv(ws, "H")
    wb.save(AUD / "pre-qual-checklist.xlsx")


def main():
    gen_offering_pptx(); gen_qual_docx(); gen_discovery_pptx(); gen_discovery_xlsx()
    gen_waf_xlsx(); gen_assess_xlsx(); gen_dod_docx()
    gen_hld_docx(); gen_lld_docx(); gen_runbook_docx(); gen_kt_docx(); gen_hyper_docx()
    gen_evidence_xlsx(); gen_prequal_xlsx()
    print("OK")

if __name__ == "__main__":
    main()
